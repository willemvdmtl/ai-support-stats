#!/usr/bin/env python3
"""Derive monthly team composition from Jira assignees."""

import html
import json
import os
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple


JIRA_CREATED_DIR = "cache/jira/by-created"
JIRA_MANIFEST = "cache/jira/manifest.json"
OUTPUT_DIR = "reports/team-composition"
CORRECTIONS_FILE = "config/team-composition-corrections.json"


def load_corrections(path: str = CORRECTIONS_FILE) -> dict:
    """Load optional manual corrections config for team composition outputs."""
    if not os.path.exists(path):
        return {}
    with open(path) as f:
        data = json.load(f)
    return data if isinstance(data, dict) else {}


def _months_in_range(months: List[str], start: str, end: str) -> List[str]:
    return [month for month in months if start <= month <= end]


def _apply_assignment_timeline_corrections(
    months: List[str],
    timeline: List[dict],
    corrections: Optional[dict],
) -> List[dict]:
    """Apply manual corrections to assignment-based timeline entries."""
    if not corrections:
        return timeline

    global_excluded = set(corrections.get("exclude_names") or [])
    assignment = corrections.get("assignment_timeline") or {}
    include_names = set(assignment.get("include_names") or [])
    overrides = assignment.get("overrides") or {}

    corrected = []
    for entry in timeline:
        name = entry.get("displayName", "")
        if name in global_excluded and name not in include_names:
            continue

        override = overrides.get(name) or {}
        if override.get("exclude"):
            continue

        active = sorted(set(entry.get("months_active") or []))
        if not active:
            continue

        if override.get("replace_months_active"):
            active = sorted(
                month for month in set(override.get("replace_months_active") or []) if month in months
            )

        if override.get("continuous_from_first_to_last") and active:
            active = _months_in_range(months, active[0], active[-1])

        if override.get("active_ranges"):
            ranged = set()
            for pair in override.get("active_ranges") or []:
                if not isinstance(pair, list) or len(pair) != 2:
                    continue
                start, end = pair
                ranged.update(_months_in_range(months, str(start), str(end)))
            if ranged:
                active = sorted(ranged)

        active_from = override.get("active_from")
        if active_from:
            active = [month for month in active if month >= str(active_from)]

        active_to = override.get("active_to")
        if active_to:
            active = [month for month in active if month <= str(active_to)]

        if not active:
            continue

        new_entry = dict(entry)
        new_entry["months_active"] = active
        month_event_counts = {
            month: int(count)
            for month, count in (entry.get("month_event_counts") or {}).items()
            if month in set(active)
        }
        new_entry["month_event_counts"] = month_event_counts
        new_entry["first_seen"] = active[0]
        new_entry["last_seen"] = active[-1]
        corrected.append(new_entry)

    return sorted(corrected, key=lambda x: (x["first_seen"], x["displayName"]))


def manifest_months() -> List[str]:
    """Return months present in the Jira manifest, sorted ascending."""
    if not os.path.exists(JIRA_MANIFEST):
        return []
    with open(JIRA_MANIFEST) as f:
        data = json.load(f)
    return sorted(data.get("months", {}).keys())


def extract_assignees(month: str) -> Optional[Dict[str, dict]]:
    """
    Extract unique assignees from the by-created cache for a given month.
    Returns a dict keyed by accountId, or None if the cache file doesn't exist.
    """
    path = os.path.join(JIRA_CREATED_DIR, f"issues-{month}.json")
    if not os.path.exists(path):
        return None
    with open(path) as f:
        data = json.load(f)
    assignees = {}
    for issue in data.get("issues", []):
        a = issue.get("fields", {}).get("assignee")
        if a and a.get("accountId"):
            aid = a["accountId"]
            if aid not in assignees:
                assignees[aid] = {
                    "accountId": aid,
                    "displayName": a.get("displayName", ""),
                    "emailAddress": a.get("emailAddress", ""),
                }
    return assignees


def load_saved(month: str) -> Optional[Dict[str, dict]]:
    """Load a previously written team-composition file, keyed by accountId."""
    path = os.path.join(OUTPUT_DIR, f"assignees-{month}.json")
    if not os.path.exists(path):
        return None
    with open(path) as f:
        data = json.load(f)
    return {a["accountId"]: a for a in data.get("assignees", [])}


def write_month(month: str, assignees: Dict[str, dict]) -> str:
    """Write the assignees JSON for a month. Returns the output path."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, f"assignees-{month}.json")
    result = {
        "month": month,
        "assignee_count": len(assignees),
        "assignees": sorted(assignees.values(), key=lambda x: x["displayName"]),
    }
    with open(path, "w") as f:
        json.dump(result, f, indent=2)
    return path


def diff(
    prev: Optional[Dict[str, dict]], curr: Dict[str, dict]
) -> Tuple[List[dict], List[dict]]:
    """Return (new_members, gone_members) by comparing prev and curr dicts."""
    if prev is None:
        return [], []
    prev_ids = set(prev.keys())
    curr_ids = set(curr.keys())
    new = sorted(
        [curr[i] for i in curr_ids - prev_ids], key=lambda x: x["displayName"]
    )
    gone = sorted(
        [prev[i] for i in prev_ids - curr_ids], key=lambda x: x["displayName"]
    )
    return new, gone


def build_timeline(months: List[str]) -> List[dict]:
    """
    Build a per-person timeline from saved monthly snapshot files.

    For each person tracks:
      - first_seen: earliest month they appear in any snapshot
      - last_seen:  latest month they appear in any snapshot
      - months_active: sorted list of months they appear in

    Only considers months that have a saved snapshot file.
    """
    # person_id -> {meta, months_active set}
    people: Dict[str, dict] = {}

    for month in months:
        saved = load_saved(month)
        if saved is None:
            continue
        for aid, info in saved.items():
            if aid not in people:
                people[aid] = {
                    "accountId": aid,
                    "displayName": info["displayName"],
                    "emailAddress": info["emailAddress"],
                    "months_active": set(),
                }
            people[aid]["months_active"].add(month)

    timeline = []
    for entry in people.values():
        active = sorted(entry["months_active"])
        timeline.append(
            {
                "accountId": entry["accountId"],
                "displayName": entry["displayName"],
                "emailAddress": entry["emailAddress"],
                "first_seen": active[0],
                "last_seen": active[-1],
                "months_active": active,
            }
        )

    return sorted(timeline, key=lambda x: (x["first_seen"], x["displayName"]))


def write_timeline(months: List[str], timeline: List[dict]) -> str:
    """Write the consolidated timeline JSON. Returns the output path."""
    import datetime as dt

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, "timeline.json")
    result = {
        "generated_at": dt.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "range": {"from": months[0], "to": months[-1]},
        "member_count": len(timeline),
        "members": timeline,
    }
    with open(path, "w") as f:
        json.dump(result, f, indent=2)
    return path


def _parse_jira_datetime(raw: str) -> datetime:
    normalized = re.sub(r"([+-]\d{2})(\d{2})$", r"\1:\2", raw)
    normalized = normalized.replace("Z", "+00:00")
    return datetime.fromisoformat(normalized)


def build_assignment_timeline(
    months: List[str],
    exclude_names: Optional[set] = None,
    min_events: int = 5,
    corrections: Optional[dict] = None,
) -> List[dict]:
    """
    Build per-person active months from assignee changelog events.

    A month is active for a person if at least one assignee-change event assigns
    a ticket to that person in that month.
    """
    exclude_names = exclude_names or set()
    people: Dict[str, dict] = {}

    observed_months = set(months)
    for filename in os.listdir(JIRA_CREATED_DIR):
        if not (filename.startswith("issues-") and filename.endswith(".json")):
            continue
        path = os.path.join(JIRA_CREATED_DIR, filename)
        with open(path) as f:
            data = json.load(f)

        for issue in data.get("issues", []):
            histories = issue.get("changelog", {}).get("histories", [])
            for history in histories:
                changed_at = history.get("created")
                if not changed_at:
                    continue
                month_key = _parse_jira_datetime(changed_at).strftime("%Y-%m")
                if month_key not in observed_months:
                    continue

                for item in history.get("items", []):
                    if item.get("field") != "assignee":
                        continue
                    name = item.get("toString")
                    if not name or name in exclude_names:
                        continue
                    if name not in people:
                        people[name] = {
                            "displayName": name,
                            "months_active": set(),
                            "events": 0,
                            "month_event_counts": {},
                        }
                    people[name]["months_active"].add(month_key)
                    people[name]["events"] += 1
                    people[name]["month_event_counts"][month_key] = (
                        people[name]["month_event_counts"].get(month_key, 0) + 1
                    )

    assignment_cfg = (corrections or {}).get("assignment_timeline") or {}
    include_names = set(assignment_cfg.get("include_names") or [])

    timeline = []
    for person in people.values():
        if person["events"] < min_events and person["displayName"] not in include_names:
            continue
        active = sorted(person["months_active"])
        timeline.append(
            {
                "displayName": person["displayName"],
                "first_seen": active[0],
                "last_seen": active[-1],
                "months_active": active,
                "events": person["events"],
                "month_event_counts": {
                    month: int(count)
                    for month, count in (person.get("month_event_counts") or {}).items()
                    if month in set(active)
                },
            }
        )

    timeline = sorted(timeline, key=lambda x: (x["first_seen"], x["displayName"]))
    return _apply_assignment_timeline_corrections(months, timeline, corrections)


def write_timeline_svg(
    months: List[str],
    timeline: List[dict],
    title: str = "Team Composition Timeline",
    subtitle: Optional[str] = None,
    monthly_rows: Optional[List[dict]] = None,
) -> str:
    """Write an SVG grid showing which months each engineer was active."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, "timeline.svg")

    subtitle_y = 50
    month_label_y = 96
    month_headcount_y = 112
    left_margin = 280
    top_margin = 118
    right_margin = 40
    bottom_margin = 50
    cell_width = 28
    cell_height = 22
    row_gap = 6
    name_x = 18

    width = left_margin + (len(months) * cell_width) + right_margin
    height = top_margin + (len(timeline) * (cell_height + row_gap)) + bottom_margin

    lines = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        (
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" '
            f'height="{height}" viewBox="0 0 {width} {height}">'
        ),
        "<style>",
        ".bg { fill: #f6f7f4; }",
        ".title { fill: #162521; font: 700 24px Helvetica, Arial, sans-serif; }",
        ".subtitle { fill: #51605a; font: 12px Helvetica, Arial, sans-serif; }",
        ".name { fill: #1f312c; font: 12px Helvetica, Arial, sans-serif; }",
        ".month { fill: #48615a; font: 11px Helvetica, Arial, sans-serif; }",
        ".grid { fill: #edf1ee; stroke: #d6ddd8; stroke-width: 1; }",
        ".active { fill: #0f766e; stroke: #0b5b54; stroke-width: 1; }",
        ".rowline { stroke: #e1e7e3; stroke-width: 1; }",
        ".legend { fill: #33443e; font: 12px Helvetica, Arial, sans-serif; }",
        "</style>",
        f'<rect class="bg" x="0" y="0" width="{width}" height="{height}" />',
        f'<text class="title" x="18" y="34">{html.escape(title)}</text>',
        (
            f'<text class="subtitle" x="18" y="{subtitle_y}">'
            f'{html.escape(subtitle or f"Active assignees by month from {months[0]} to {months[-1]}")}'
            "</text>"
        ),
    ]
    # Build headcount lookup
    headcount_by_month = {}
    if monthly_rows:
        for row in monthly_rows:
            headcount_by_month[row["month"]] = row["headcount"]

    for index, month in enumerate(months):
        x = left_margin + (index * cell_width) + (cell_width / 2)
        label = html.escape(month)
        lines.append(
            (
                f'<text class="month" x="{x:.1f}" y="{month_label_y}" '
                f'text-anchor="start" transform="rotate(-45 {x:.1f} {month_label_y})">{label}</text>'
            )
        )
        # Add headcount as secondary label
        if month in headcount_by_month:
            headcount = headcount_by_month[month]
            lines.append(
                (
                    f'<text class="month" x="{x:.1f}" y="{month_headcount_y}" '
                    f'text-anchor="middle" font-size="10">{headcount}</text>'
                )
            )

    for row_index, member in enumerate(timeline):
        y = top_margin + (row_index * (cell_height + row_gap))
        name = html.escape(member["displayName"])
        active_months = set(member["months_active"])

        lines.append(
            f'<line class="rowline" x1="18" y1="{y - 4}" x2="{width - 20}" y2="{y - 4}" />'
        )
        lines.append(
            f'<text class="name" x="{name_x}" y="{y + 15}">{name}</text>'
        )

        for col_index, month in enumerate(months):
            x = left_margin + (col_index * cell_width)
            klass = "active" if month in active_months else "grid"
            lines.append(
                f'<rect class="{klass}" x="{x}" y="{y}" width="20" height="20" rx="3" />'
            )

    lines.append("</svg>")

    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    return path


def _next_month_key(month: str) -> str:
    year = int(month[:4])
    month_number = int(month[5:7])
    if month_number == 12:
        return f"{year + 1}-01"
    return f"{year}-{month_number + 1:02d}"


def build_monthly_roster_from_summary(summary: dict) -> List[dict]:
    """Build monthly roster rows from the inactivity-based membership summary."""
    first_month = summary.get("first_month")
    if not first_month:
        return []

    def next_month_key(month: str) -> str:
        year = int(month[:4])
        month_number = int(month[5:7])
        if month_number == 12:
            return f"{year + 1}-01"
        return f"{year}-{month_number + 1:02d}"

    events_by_month = {
        event.get("month"): {
            "joins": list(event.get("joins") or []),
            "leaves": list(event.get("leaves") or []),
        }
        for event in (summary.get("events") or [])
        if event.get("month")
    }

    months = [first_month]
    end_month = first_month
    range_info = summary.get("range") or {}
    range_to = range_info.get("to")
    if range_to:
        end_month = str(range_to)
    if events_by_month:
        end_month = max(end_month, max(events_by_month.keys()))
    while months[-1] < end_month:
        months.append(next_month_key(months[-1]))

    rows = []
    roster = set(summary.get("names") or [])
    for index, month in enumerate(months):
        event = events_by_month.get(month) or {"joins": [], "leaves": []}
        if index > 0:
            roster.update(event["joins"])
            roster.difference_update(event["leaves"])
        rows.append(
            {
                "month": month,
                "headcount": len(roster),
                "joined": list(event["joins"] if index > 0 else []),
                "left": list(event["leaves"] if index > 0 else []),
                "roster": sorted(roster),
            }
        )

    return rows


def write_assignment_timeline_excel(
    months: List[str],
    timeline: List[dict],
    output_path: Optional[str] = None,
) -> str:
    """Write Excel workbook from assignment timeline data used by the SVG visual."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Missing dependency: openpyxl")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = output_path or os.path.join(OUTPUT_DIR, "team-composition.xlsx")

    corrections = load_corrections()
    summary_months = list(months)
    summary_months.append(_next_month_key(months[-1]))
    summary = build_membership_summary(
        summary_months,
        exclude_names={name for name in (corrections.get("exclude_names") or [])},
        inactivity_threshold=3,
        corrections=corrections,
    )
    monthly_rows = build_monthly_roster_from_summary(summary)
    members = sorted(timeline, key=lambda entry: (entry["first_seen"], entry["displayName"]))
    roster_by_month = {row["month"]: set(row["roster"]) for row in monthly_rows}

    thin = Side(style="thin", color="C7D7CE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1B4332")
    alt_fill = PatternFill("solid", fgColor="F0F4F1")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    join_fill = PatternFill("solid", fgColor="D1FAE5")
    left_fill = PatternFill("solid", fgColor="FEE2E2")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top")

    wb = openpyxl.Workbook()

    # Sheet 1: Monthly roster changes
    ws1 = wb.active
    ws1.title = "Assignment Monthly"
    ws1.freeze_panes = "A2"
    cols1 = [
        ("Month", 12),
        ("Headcount", 12),
        ("Joined", 35),
        ("Left", 35),
        ("Roster", 65),
    ]
    for col, (label, width) in enumerate(cols1, 1):
        cell = ws1.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws1.column_dimensions[get_column_letter(col)].width = width

    for row_index, row in enumerate(monthly_rows, 2):
        has_joins = bool(row["joined"])
        has_left = bool(row["left"])
        base_fill = white_fill if row_index % 2 == 0 else alt_fill
        values = [
            row["month"],
            row["headcount"],
            ", ".join(row["joined"]),
            ", ".join(row["left"]),
            ", ".join(row["roster"]),
        ]
        for col, value in enumerate(values, 1):
            cell = ws1.cell(row=row_index, column=col, value=value)
            cell.border = border
            cell.alignment = wrap if col in (3, 4, 5) else center
            if col == 3 and has_joins:
                cell.fill = join_fill
            elif col == 4 and has_left:
                cell.fill = left_fill
            else:
                cell.fill = base_fill

    # Sheet 2: Member assignment spans
    ws2 = wb.create_sheet("Assignment Members")
    ws2.freeze_panes = "A2"
    cols2 = [
        ("Name", 30),
        ("First Active Month", 18),
        ("Last Active Month", 18),
        ("Active Months", 14),
        ("Assignment Events", 18),
    ]
    for col, (label, width) in enumerate(cols2, 1):
        cell = ws2.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws2.column_dimensions[get_column_letter(col)].width = width

    member_rows = []
    for member in members:
        active_months = [
            month for month in months if member["displayName"] in roster_by_month.get(month, set())
        ]
        if not active_months:
            continue
        member_rows.append((member, active_months))

    for row_index, (member, active_months) in enumerate(member_rows, 2):
        base_fill = white_fill if row_index % 2 == 0 else alt_fill
        values = [
            member["displayName"],
            active_months[0],
            active_months[-1],
            len(active_months),
            member.get("events", ""),
        ]
        for col, value in enumerate(values, 1):
            cell = ws2.cell(row=row_index, column=col, value=value)
            cell.border = border
            cell.alignment = wrap if col == 1 else center
            cell.fill = base_fill

    # Sheet 3: Monthly headcount (fixed vs floating)
    team_cfg = corrections.get("team_assignments") or {}
    floating_set = set(team_cfg.get("floating") or [])

    ws3 = wb.create_sheet("Headcount (Fixed vs Floating)")
    ws3.freeze_panes = "A2"
    cols3 = [
        ("Month", 12),
        ("Fixed", 10),
        ("Floating", 12),
        ("Total", 10),
        ("Fixed Members", 50),
        ("Floating Members", 50),
    ]
    for col, (label, width) in enumerate(cols3, 1):
        cell = ws3.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws3.column_dimensions[get_column_letter(col)].width = width

    for row_index, row in enumerate(monthly_rows, 2):
        base_fill = white_fill if row_index % 2 == 0 else alt_fill
        roster = row["roster"]
        fixed_members = [name for name in roster if name not in floating_set]
        floating_members = [name for name in roster if name in floating_set]
        fixed_count = len(fixed_members)
        floating_count = len(floating_members)
        total_count = len(roster)

        values = [
            row["month"],
            fixed_count,
            floating_count,
            total_count,
            ", ".join(sorted(fixed_members)) if fixed_members else "",
            ", ".join(sorted(floating_members)) if floating_members else "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws3.cell(row=row_index, column=col, value=value)
            cell.border = border
            cell.alignment = wrap if col in (5, 6) else center
            cell.fill = base_fill

    wb.save(path)
    return path


def write_team_split_excel(
    months: List[str],
    output_path: Optional[str] = None,
) -> str:
    """Write a separate Excel file with fixed vs floating headcount breakdown by month."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Missing dependency: openpyxl")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = output_path or os.path.join(OUTPUT_DIR, "team-composition-team-split.xlsx")

    corrections = load_corrections()
    summary_months = list(months)
    summary_months.append(_next_month_key(months[-1]))
    summary = build_membership_summary(
        summary_months,
        exclude_names={name for name in (corrections.get("exclude_names") or [])},
        inactivity_threshold=3,
        corrections=corrections,
    )
    monthly_rows = build_monthly_roster_from_summary(summary)

    team_cfg = corrections.get("team_assignments") or {}
    floating_set = set(team_cfg.get("floating") or [])

    thin = Side(style="thin", color="C7D7CE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1B4332")
    alt_fill = PatternFill("solid", fgColor="F0F4F1")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Headcount & Shared"
    ws.freeze_panes = "A2"

    join_fill = PatternFill("solid", fgColor="D1FAE5")
    left_fill = PatternFill("solid", fgColor="FEE2E2")

    cols = [
        ("Month", 12),
        ("Headcount", 12),
        ("Shared", 12),
        ("Joins", 35),
        ("Leaves", 35),
        ("Headcount Members", 50),
        ("Shared Members", 50),
    ]
    for col, (label, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width

    for row_index, row in enumerate(monthly_rows, 2):
        base_fill = white_fill if row_index % 2 == 0 else alt_fill
        roster = row["roster"]
        fixed_members = [name for name in roster if name not in floating_set]
        floating_members = [name for name in roster if name in floating_set]
        fixed_count = len(fixed_members)
        floating_count = len(floating_members)
        joins = ", ".join(row.get("joined") or [])
        leaves = ", ".join(row.get("left") or [])
        has_joins = bool(row.get("joined"))
        has_left = bool(row.get("left"))

        values = [
            row["month"],
            fixed_count,
            floating_count,
            joins,
            leaves,
            ", ".join(sorted(fixed_members)) if fixed_members else "",
            ", ".join(sorted(floating_members)) if floating_members else "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_index, column=col, value=value)
            cell.border = border
            cell.alignment = wrap if col in (4, 5, 6, 7) else center
            if col == 4 and has_joins:
                cell.fill = join_fill
            elif col == 5 and has_left:
                cell.fill = left_fill
            else:
                cell.fill = base_fill

    wb.save(path)
    return path


def build_membership_summary(
    months: List[str],
    exclude_names: Optional[set] = None,
    inactivity_threshold: int = 3,
    corrections: Optional[dict] = None,
) -> dict:
    """
    Build a month-by-month membership summary from assignment changelog activity.

    A leave is confirmed once inactivity exceeds the threshold, but the leave is
    recorded against the first inactive month after the member's last active month.
    """
    exclude_names = exclude_names or set()
    corrections = corrections or {}
    assignment_cfg = corrections.get("assignment_timeline") or {}
    overrides = assignment_cfg.get("overrides") or {}
    min_events = int(assignment_cfg.get("min_events") or 5)
    join_confirmation_cfg = assignment_cfg.get("join_confirmation") or {}

    assignment_timeline = build_assignment_timeline(
        months,
        exclude_names=exclude_names,
        min_events=min_events,
        corrections=corrections,
    )

    by_name = {entry.get("displayName"): entry for entry in assignment_timeline if entry.get("displayName")}
    month_index = {month: index for index, month in enumerate(months)}

    def _join_rule_for(name: str) -> dict:
        person_cfg = ((join_confirmation_cfg.get("overrides") or {}).get(name) or {})
        window_months = int(
            person_cfg.get(
                "window_months",
                join_confirmation_cfg.get("window_months", 3),
            )
        )
        min_active_months = int(
            person_cfg.get(
                "min_active_months_in_window",
                join_confirmation_cfg.get("min_active_months_in_window", 1),
            )
        )
        min_events_in_join_month = int(
            person_cfg.get(
                "min_events_in_join_month",
                join_confirmation_cfg.get("min_events_in_join_month", 1),
            )
        )
        return {
            "window_months": max(1, window_months),
            "min_active_months_in_window": max(1, min_active_months),
            "min_events_in_join_month": max(1, min_events_in_join_month),
        }

    def _passes_join_confirmation(name: str, month: str) -> bool:
        person = by_name.get(name) or {}
        month_event_counts = person.get("month_event_counts") or {}
        month_active = set(person.get("months_active") or [])
        if not month_active:
            return False

        rule = _join_rule_for(name)
        if month_event_counts.get(month, 0) < rule["min_events_in_join_month"]:
            return False

        start_index = month_index.get(month)
        if start_index is None:
            return False
        end_index = min(len(months), start_index + rule["window_months"])
        window = months[start_index:end_index]
        active_in_window = sum(1 for value in window if value in month_active)
        return active_in_window >= rule["min_active_months_in_window"]

    active_by_month: Dict[str, set] = {month: set() for month in months}
    for member in assignment_timeline:
        name = member.get("displayName")
        if not name:
            continue
        for month in member.get("months_active") or []:
            if month in active_by_month:
                active_by_month[month].add(name)

    if not months:
        return {"first_month": None, "headcount": 0, "names": [], "events": []}

    status: Dict[str, dict] = {}
    pending_events: Dict[str, dict] = {}

    first_month = months[0]
    initial = active_by_month[first_month]
    for name in initial:
        status[name] = {
            "member": True,
            "inactive_streak": 0,
            "first_inactive_month": None,
        }

    initial_names = sorted(initial)

    # Expand continuous overrides across the observed span so short gaps don't
    # fragment a single stint into leave/join pairs.
    for name, override in overrides.items():
        if not override.get("continuous_from_first_to_last"):
            continue
        observed_months = [month for month in months if name in active_by_month[month]]
        if len(observed_months) < 2:
            continue
        start_month = observed_months[0]
        end_month = observed_months[-1]
        for month in months:
            if start_month <= month <= end_month:
                active_by_month[month].add(name)

    def sort_names(values: List[str]) -> List[str]:
        return sorted(str(value) for value in values if value)

    for month in months[1:]:
        month_active = active_by_month[month]
        known_names = set(status.keys()) | set(month_active)

        for name in sorted(known_names):
            if name not in status:
                status[name] = {
                    "member": False,
                    "inactive_streak": 0,
                    "first_inactive_month": None,
                }

            person = status[name]
            is_active = name in month_active

            if person["member"]:
                if is_active:
                    person["inactive_streak"] = 0
                    person["first_inactive_month"] = None
                else:
                    person["inactive_streak"] += 1
                    if person["inactive_streak"] == 1:
                        person["first_inactive_month"] = month
                    if person["inactive_streak"] > inactivity_threshold:
                        person["member"] = False
                        leave_month = person["first_inactive_month"] or month
                        event = pending_events.setdefault(
                            leave_month,
                            {"month": leave_month, "joins": [], "leaves": []},
                        )
                        event["leaves"].append(name)
            else:
                if is_active:
                    if _passes_join_confirmation(name, month):
                        person["member"] = True
                        person["inactive_streak"] = 0
                        person["first_inactive_month"] = None
                        event = pending_events.setdefault(
                            month,
                            {"month": month, "joins": [], "leaves": []},
                        )
                        event["joins"].append(name)

    headcount = len(initial_names)
    events = []
    for month in months[1:]:
        event = pending_events.get(month)
        if not event or (not event["joins"] and not event["leaves"]):
            continue
        event["joins"] = sort_names(event["joins"])
        event["leaves"] = sort_names(event["leaves"])
        headcount += len(event["joins"]) - len(event["leaves"])
        event["headcount"] = headcount
        events.append(event)

    return {
        "first_month": first_month,
        "headcount": len(initial_names),
        "names": initial_names,
        "events": events,
        "range": {"from": months[0], "to": months[-1]},
    }
